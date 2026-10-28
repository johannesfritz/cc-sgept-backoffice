"""A standard deliverable invoice keeps each amount in the Amount column, whatever the description length (JCC-9404)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Twips

RECIPIENT = {"company": "Test Client AG", "name": "Jane Doe", "street": "Street 1",
             "city": "8000 Zurich", "country": "Switzerland"}
LONG = "Policy intelligence dashboard licence, Q4 2026 XYZ"
VERY_LONG = ("Policy intelligence dashboard licence, Q4 2026, all users of the Ministry of Trade "
             "and its agencies & partners, renewed.")


def generate(generator, items, items_type, total):
    generator.generate_standard_invoice(
        invoice_number="99999", invoice_date="28 October 2026", currency="CHF", recipient=RECIPIENT,
        subject_line="Render check", intro_paragraph="A render check.", items=items,
        items_type=items_type, total=total, vat_note="No VAT is applied.", sync=False)
    (path,) = generator.OUTPUT_DIR.glob("*/*.docx")
    return Document(path)


def right_tabs(paragraph):
    return [(stop.position, stop.alignment) for stop in paragraph.paragraph_format.tab_stops]


def test_header_and_total_share_one_right_tab(generator):
    doc = generate(generator, [{"description": "Short licence fee", "amount": 1200}], "deliverable", 1200)
    rows = {p.text.split("\t")[0]: p for p in doc.paragraphs if "\t" in p.text}
    for label, amount in {"Description": "Amount (CHF)", "Invoice total:": "1,200.00"}.items():
        assert rows[label].text == f"{label}\t{amount}"
        assert right_tabs(rows[label]) == [(Twips(9026), WD_TAB_ALIGNMENT.RIGHT)]


def test_long_descriptions_wrap_inside_their_own_column(generator):
    assert len(LONG) == 50 and len(VERY_LONG) == 120
    items = [{"description": d, "amount": a} for d, a in
             [("Short licence fee", 1200), (LONG, 12000), (VERY_LONG, -500)]]
    doc = generate(generator, items, "deliverable", 12700)
    (table,) = doc.tables
    assert [[c.text for c in row.cells] for row in table.rows] == [
        ["Short licence fee", "1,200.00"], [LONG, "12,000.00"], [VERY_LONG, "-500.00"]]
    tbl = table._tbl
    # The table runs from the Indented text edge (2835) to the amount tab (9026).
    tbl_pr = tbl.find(qn("w:tblPr"))
    assert tbl_pr.find(qn("w:tblInd")).get(qn("w:w")) == "2835"
    # Zero cell margins override the template's default Normal Table (108 twips each side).
    margins = tbl_pr.find(qn("w:tblCellMar"))
    assert [margins.find(qn(f"w:{side}")).get(qn("w:w")) for side in ("left", "right")] == ["0", "0"]
    assert [int(c.get(qn("w:w"))) for c in tbl.tblGrid.findall(qn("w:gridCol"))] == [4717, 1474]
    for row in table.rows:
        description, amount = row.cells
        assert description.paragraphs[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT  # not justified
        assert amount.paragraphs[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT
        assert description.paragraphs[0].paragraph_format.left_indent == 0
    # The item rows sit between the header and the total.
    body = list(doc.element.body)
    header = next(i for i, p in enumerate(body) if p.tag == qn("w:p")
                  and "".join(p.itertext()).startswith("Description"))
    assert body[header + 1] is tbl
    assert "".join(body[header + 2].itertext()).startswith("Invoice total:")


def test_hours_invoice_keeps_its_columns(generator):
    doc = generate(generator, [{"description": "J. Fritz", "hours": 45, "rate": 300, "total": 13500}],
                   "hours", 13500)
    paragraphs = doc.paragraphs
    assert doc.tables == []
    texts = [p.text for p in paragraphs]
    assert "Description\t\t\tHours\tRate (CHF)\t\tTotal (CHF)" in texts
    assert "J. Fritz\t\t\t45\t300.00\t\t13,500.00" in texts
    assert "Invoice total:\t\t\t\t\t\t13,500.00" in texts
    assert all(right_tabs(p) == [] for p in paragraphs if p.text.startswith(("Description", "J. Fritz", "Invoice total")))
